import os
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "..", "data")
OUT = os.path.join(HERE, "..", "output")
os.makedirs(OUT, exist_ok=True)


def read_grid(path, sheet="Sheet1"):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    grid = {}
    for row in ws.iter_rows():
        for c in row:
            grid[(c.row, c.column)] = c.value
    return grid, ws.max_row


def find_blocks(grid, max_row, row_letters):
    blocks = []
    r = 1
    while r <= max_row:
        v = grid.get((r, 1))
        if isinstance(v, str) and v not in row_letters:
            title = v
            header_row = r + 1
            data_rows = []
            rr = header_row + 1
            while rr <= max_row and grid.get((rr, 1)) in row_letters:
                data_rows.append(rr)
                rr += 1
            blocks.append({"title": title, "data_rows": data_rows})
            r = rr
        else:
            r += 1
    return blocks


def block_to_long(grid, block, value_name):
    recs = []
    for r in block["data_rows"]:
        row_letter = grid[(r, 1)]
        for col in range(2, 14):  
            v = grid.get((r, col))
            if v is None:
                continue
            well_col = col - 1
            recs.append({"row": row_letter, "col": well_col,
                         "well": f"{row_letter}{well_col:02d}", value_name: v})
    return pd.DataFrame(recs)


def parse_file(path, row_letters):
    grid, max_row = read_grid(path)
    blocks = find_blocks(grid, max_row, row_letters)
    assert len(blocks) == 4, f"Expected 4 data blocks, found {len(blocks)}: {[b['title'] for b in blocks]}"
    platemap_b, gfp_b, mch_b, od_b = blocks
    platemap = block_to_long(grid, platemap_b, "label")
    gfp = block_to_long(grid, gfp_b, "GFP_raw")
    mch = block_to_long(grid, mch_b, "mCherry_raw")
    od = block_to_long(grid, od_b, "OD600")

    df = platemap.merge(gfp, on=["row", "col", "well"], how="outer")
    df = df.merge(mch, on=["row", "col", "well"], how="outer")
    df = df.merge(od, on=["row", "col", "well"], how="outer")
    df["label"] = df["label"].astype(str)
    return df, [b["title"] for b in blocks]


if __name__ == "__main__":
    df920, titles920 = parse_file(os.path.join(DATA, "20260920.xlsx"), set("ABDE"))
    print("0920 blocks:", titles920)
    df920.to_csv(os.path.join(OUT, "df920_tidy.csv"), index=False)

    df921, titles921 = parse_file(os.path.join(DATA, "20260921.xlsx"), set("ABCDEFG"))
    print("0921 blocks:", titles921)
    df921.to_csv(os.path.join(OUT, "df921_tidy.csv"), index=False)

    print(f"\n0920 rows: {len(df920)}   0921 rows: {len(df921)}")
    print("\n0920 label counts:\n", df920["label"].value_counts())
    print("\n0921 label counts:\n", df921["label"].value_counts())
